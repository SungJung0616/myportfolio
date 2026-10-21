"""Build review-only DOCX/PDF from one source. Never replace the public resume."""
from pathlib import Path
import json
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import letter

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'resume/systems-resume-content.json').read_text(encoding='utf-8'))
OUT = ROOT / 'resume/Sung_Jung_Business_Systems_Resume'


def blocks():
    yield 'title', DATA['name']
    yield 'headline', DATA['headline']
    yield 'contact', DATA['contact']
    yield 'links', DATA['links']
    yield 'heading', 'Professional Summary'
    yield 'body', DATA['summary']
    yield 'heading', 'Core Competencies'
    yield 'body', DATA['competencies']
    yield 'heading', 'Professional Experience'
    for role in DATA['roles']:
        if role.get('pageBreak'):
            yield 'break', ''
            yield 'heading', 'Professional Experience Continued'
        yield 'role', f"{role['title']} | {role['employer']}"
        yield 'meta', f"{role['location']} | {role['dates']}"
        for item in role['bullets']:
            yield 'bullet', item
    yield 'heading', 'Selected Projects'
    for project in DATA['projects']:
        yield 'project', project
        yield 'body', project['text']
    yield 'heading', 'Technical Skills'
    for line in DATA['skills']:
        yield 'body', line
    yield 'heading', 'Education and Certifications'
    yield 'body', DATA['education']
    yield 'body', DATA['certifications']


def add_link(paragraph, label, url):
    rid = paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rid)
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = label
    run.append(text); link.append(run); paragraph._p.append(link)


def docx():
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.55)
    sec.left_margin = sec.right_margin = Inches(.65)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]; style.font.name = 'Arial'; style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(10.2)
    doc.styles['Normal'].paragraph_format.space_after = Pt(3)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.05
    for kind, value in blocks():
        if kind == 'break': doc.add_page_break(); continue
        if kind == 'links':
            p = doc.add_paragraph()
            for index, (label, url) in enumerate(value):
                if index: p.add_run(' | ')
                add_link(p, label, url)
            continue
        if kind == 'project':
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            add_link(p, value['name'], value['url'])
            continue
        p = doc.add_paragraph(style='Title' if kind == 'title' else 'Heading 1' if kind == 'heading' else 'Normal')
        run = p.add_run(('- ' if kind == 'bullet' else '') + value)
        if kind == 'title': run.font.size = Pt(22); p.paragraph_format.space_after = Pt(3)
        if kind == 'heading':
            run.font.size = Pt(11); run.bold = True
            p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4)
        if kind in ['role', 'headline']: run.bold = True
        if kind == 'meta': run.font.size = Pt(9.4)
        if kind in ['title', 'headline', 'role', 'meta', 'heading']: p.paragraph_format.keep_with_next = True
        if kind == 'bullet': p.paragraph_format.left_indent = Inches(.1); p.paragraph_format.first_line_indent = Inches(-.1)
    doc.core_properties.title = 'Sung Jung Business Systems Resume'
    doc.core_properties.author = 'Sung Jung'
    doc.save(str(OUT) + '.docx')


def pdf():
    styles = {key: ParagraphStyle(key, fontName='Helvetica', fontSize=10.2, leading=12.4, spaceAfter=3) for key in ['title','headline','contact','body','heading','role','meta','bullet','links','project']}
    styles['title'].fontSize = 22; styles['title'].leading = 25
    for key in ['title','headline','heading','role','project']: styles[key].fontName = 'Helvetica-Bold'; styles[key].keepWithNext = True
    styles['heading'].fontSize = 11; styles['heading'].spaceBefore = 8; styles['heading'].spaceAfter = 4
    styles['meta'].fontSize = 9.4; styles['meta'].keepWithNext = True
    styles['bullet'].leftIndent = 8; styles['bullet'].firstLineIndent = -8
    story = []
    for kind, value in blocks():
        if kind == 'break': story.append(PageBreak()); continue
        if kind == 'project': text = f'<a href="{escape(value["url"], quote=True)}">{escape(value["name"])}</a>'
        elif kind == 'links': text = ' | '.join(f'<a href="{escape(url, quote=True)}">{escape(label)}</a>' for label, url in value)
        else: text = ('- ' if kind == 'bullet' else '') + escape(value)
        story.append(Paragraph(text, styles[kind]))
    SimpleDocTemplate(str(OUT)+'.pdf', pagesize=letter, leftMargin=.65*72, rightMargin=.65*72, topMargin=.55*72, bottomMargin=.55*72, title='Sung Jung Business Systems Resume', author='Sung Jung').build(story)


if __name__ == '__main__':
    docx(); pdf()
    print(str(OUT) + '.docx'); print(str(OUT) + '.pdf')
