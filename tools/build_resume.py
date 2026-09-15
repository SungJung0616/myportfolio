from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "resume" / "Sung_Jung_Resume.docx"
NAVY = RGBColor(22, 42, 59)
GREEN = RGBColor(25, 126, 76)
GRAY = RGBColor(78, 88, 96)

doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Inches(.48), Inches(.45)
sec.left_margin = sec.right_margin = Inches(.58)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
normal.font.size = Pt(10.2)
normal.paragraph_format.space_after = Pt(2.2)
normal.paragraph_format.line_spacing = 1.02

def hyperlink(paragraph, text, url):
    rid = paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link, run, prop = OxmlElement('w:hyperlink'), OxmlElement('w:r'), OxmlElement('w:rPr')
    link.set(qn('r:id'), rid)
    col = OxmlElement('w:color'); col.set(qn('w:val'), '197E4C'); prop.append(col)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); prop.append(underline)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '20'); prop.append(size)
    run.append(prop)
    node = OxmlElement('w:t'); node.text = text; run.append(node)
    link.append(run); paragraph._p.append(link)

def section_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(7), Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(12.4); r.font.color.rgb = NAVY

def role(title, company, location, dates, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4.5), Pt(1.5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title); r.bold = True; r.font.size = Pt(10.7); r.font.color.rgb = NAVY
    r = p.add_run(f' | {company} | {location} | {dates}'); r.bold = True; r.font.color.rgb = GRAY
    for item in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.left_indent, bp.paragraph_format.first_line_indent = Inches(.16), Inches(-.13)
        bp.paragraph_format.space_after = Pt(1.4)
        bp.add_run(item)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
r = p.add_run('SUNG JUNG'); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
r = p.add_run('OPERATIONS SYSTEMS | QUALITY ENGINEERING | WORKFLOW AUTOMATION'); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = GREEN
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4)
p.add_run('Gardena, CA | 949-300-1641 | sungjung0616@gmail.com | ')
hyperlink(p, 'LinkedIn', 'https://www.linkedin.com/in/sungjung0616/'); p.add_run(' | ')
hyperlink(p, 'GitHub', 'https://github.com/SungJung0616'); p.add_run(' | ')
hyperlink(p, 'Portfolio', 'https://sj-personalportfolio.netlify.app/')

section_title('Professional Summary')
doc.add_paragraph('Operations systems and quality professional who improves high-volume business workflows through structured validation, practical automation, and cross-functional delivery. Experience spans 3PL operations, marketplace data transformation, PlayStation release QA, and independently delivered cloud applications. Reduced a recurring shipping-data workflow from about three hours to 30 minutes and a regression cycle from about eight hours to four while preserving human review and coverage.')

section_title('Core Skills')
skills = [
('Operations & Delivery', '3PL workflows, account coordination, process improvement, stakeholder communication, work instructions, exception management'),
('Quality & Data', 'Manual and regression testing, release validation, test cases, defect reporting, data validation, reconciliation, result review'),
('Automation & Systems', 'Python, JavaScript, Google Apps Script, Excel, Google Sheets, Extensiv WMS, Workpad OMS, Shopify, ECOUNT ERP (basic)'),
('Development', 'React, Node.js, HTML, CSS, Git/GitHub, SQL (basic), PostgreSQL, AWS configuration, Selenium projects, Unity Test Framework use')]
for label, value in skills:
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label + ': '); r.bold = True; r.font.color.rgb = NAVY; p.add_run(value)

section_title('Professional Experience')
role('Operations Manager - Account & Warehouse Operations', 'GPS Logix', 'Compton, CA', '2024 - Present', [
'Helped establish a new 3PL operation, including account-specific locations, picking and packing flows, work instructions, and daily execution across Extensiv WMS and Workpad OMS.',
'Designed an Excel-based SKU batching workflow to reduce picking travel, translated the operational requirement, and partnered with the AcrossB development team until automated SKU separation was added to Workpad.',
'Support an operation that grew from approximately 3,000-3,500 daily orders at launch to about 11,000 typical daily orders across roughly 10 client accounts; coordinate a high-volume account and surface inventory and order risks.',
'Implemented an AI-assisted Gmail, Google Sheets, and Apps Script workflow that classifies account communications, records inbound and outbound information, and distributes a daily operational summary.',
'Built a scanner-based returns workflow that uses tracking and lot inputs to populate SKU and expiration data, reducing repetitive manual entry.'])
role('QA Engineer', 'Studio Sai', 'Remote from South Korea', 'Dec 2021 - May 2024', [
'Supported development and PlayStation release quality for Eternights through gameplay, functional, regression, localization, and release-focused testing.',
'Authored reproducible defect reports, validated fixes, reviewed results, and collaborated with developers to identify regressions and clarify expected behavior.',
'Reduced a recurring regression workflow from approximately eight hours to four by using developer-provided save states and checkpoints while preserving validation coverage.',
'Used Unity Test Framework features in the existing environment for UI-focused automated validation alongside manual and regression testing.'])

doc.add_page_break()
section_title('Professional Experience Continued')
role('Database Manager', 'GPS Logix', 'Carson, CA', 'Jun 2016 - Nov 2021', [
'Analyzed a recurring Coupang order workflow and built Python/openpyxl automation to map, normalize, validate, and prepare marketplace data for the GPS shipping-system upload format.',
'Reduced approximate daily preparation time from three hours to 30 minutes while retaining human review and final system-upload control.',
'Validated customs and contact fields, normalized postal and product data, handled options and bundle quantities, and separated incomplete or suspicious records for manual review.',
'Identified recurring address and customs-document risks and helped strengthen pre-shipment validation requirements before international shipping.'])
role('Owner - Operations & Systems (Part-time)', 'LAYRD', 'Los Angeles, CA', 'Jul 2025 - Present', [
'Own and co-operate a Shopify-based home and tableware business, with responsibility for storefront systems, fulfillment support, shipping labels, returns, financial records, and business administration.',
'Configure and troubleshoot the custom domain, storefront connections, shipping setup, and customer-facing functionality; share fulfillment and take primary ownership of returns processing.',
'Maintain transaction and expense records for CPA review and independently completed the LAYRD trademark process through registration using AI-assisted research.',
'Support local sales through consignment at CLAEDO in ROW DTLA and a display, consignment, and customer-pickup partnership with Baking Room.'])

section_title('Selected Projects')
p = doc.add_paragraph(); r = p.add_run('Remembering Young Hoon - Independent Cloud Project: '); r.bold=True; r.font.color.rgb=NAVY
p.add_run('Planned and delivered a bilingual memorial archive with reviewed submissions, role-based administration, private media storage, and asynchronous AWS image processing. Independently owned feature decisions, AWS setup, validation, and deployment using AI-assisted development. '); hyperlink(p, 'Case Study', 'https://sj-personalportfolio.netlify.app/case-studies/remembering-young-hoon')
p = doc.add_paragraph(); r = p.add_run('Coupang Order Automation: '); r.bold=True; r.font.color.rgb=NAVY
p.add_run('Python/openpyxl pipeline for mapping marketplace orders, validating critical fields, flagging exceptions, and generating upload-ready shipping data. '); hyperlink(p, 'GitHub', 'https://github.com/SungJung0616')
p = doc.add_paragraph(); r = p.add_run('Career Portfolio and Case Studies: '); r.bold=True; r.font.color.rgb=NAVY
p.add_run('React and Three.js portfolio documenting verified QA, automation, operations, and end-to-end project ownership. '); hyperlink(p, 'Portfolio', 'https://sj-personalportfolio.netlify.app/')

section_title('Education & Certifications')
p = doc.add_paragraph(); r=p.add_run('Education: '); r.bold=True; r.font.color.rgb=NAVY; p.add_run('B.S. in Computer Science, California State University, Northridge')
p = doc.add_paragraph(); r=p.add_run('Certifications: '); r.bold=True; r.font.color.rgb=NAVY; p.add_run('ISTQB Certified Tester - Foundation Level; ISTQB Agile Tester; Microsoft Office Specialist')

for section in doc.sections:
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('Sung Jung | Operations Systems, Quality Engineering & Workflow Automation')
    run.font.name = 'Arial'; run.font.size = Pt(8); run.font.color.rgb = GRAY

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
